"""Small python-docx toolkit for the manuscript: numbered citations, captioned tables/figures, equations as images."""
import io
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
plt.rcParams["mathtext.fontset"] = "stix"


class Paper:
    def __init__(self, refs: dict):
        self.doc = Document()
        self.refs = refs                     # key -> formatted reference string
        self.cite_order: list[str] = []      # keys in order of first citation
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self._setup()

    # ---------------- base setup ----------------
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)       # A4
        sec.left_margin = sec.right_margin = Inches(1.0)
        sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(0.9)
        st = d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.18
        for lvl, size in ((1, 14), (2, 12), (3, 11)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = lvl == 3
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(5)
            h.paragraph_format.keep_with_next = True
        self._footer_page_numbers()

    def _footer_page_numbers(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(9)

    # ---------------- inline markup ----------------
    def _cite_text(self, keys):
        nums = []
        for k in keys:
            if k not in self.refs:
                raise KeyError(f"unknown reference key {k}")
            if k not in self.cite_order:
                self.cite_order.append(k)
            nums.append(self.cite_order.index(k) + 1)
        nums = sorted(set(nums))
        out, i = [], 0
        while i < len(nums):                      # collapse runs: 1,2,3 -> 1–3
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(str(n) for n in nums[i:j + 1]))
            i = j + 1
        return "[" + ",".join(out) + "]"

    def _runs(self, p, text, size=None, color=None):
        text = re.sub(r"\{cite:([^}]+)\}", lambda m: self._cite_text([k.strip() for k in m.group(1).split(",")]), text)
        for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^\{[^}]+\}|_\{[^}]+\})", text):
            if not tok:
                continue
            if tok.startswith("**"):
                r = p.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith("*"):
                r = p.add_run(tok[1:-1]); r.italic = True
            elif tok.startswith("^{"):
                r = p.add_run(tok[2:-1]); r.font.superscript = True
            elif tok.startswith("_{"):
                r = p.add_run(tok[2:-1]); r.font.subscript = True
            else:
                r = p.add_run(tok)
            r.font.name = FONT
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)
        return p

    # ---------------- text blocks ----------------
    def title(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(17); r.font.name = FONT

    def centered(self, text, size=11, bold=False, italic=False, space_after=4):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        self._runs(p, text, size=size)
        for r in p.runs:
            r.bold = bold or r.bold
            r.italic = italic or r.italic
        return p

    def h(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def p(self, text, justify=True, indent=False, space_after=6, keep_next=False):
        para = self.doc.add_paragraph()
        if justify:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            para.paragraph_format.first_line_indent = Inches(0.3)
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.keep_with_next = keep_next
        return self._runs(para, text)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.doc.add_paragraph(style=style)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(3)
            self._runs(para, it)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- equations ----------------
    def equation(self, latex, fontsize=13):
        self.eq_n += 1
        fig = plt.figure(figsize=(0.1, 0.1))
        fig.text(0, 0, f"${latex}$", fontsize=fontsize)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=330, bbox_inches="tight", pad_inches=0.04, transparent=True)
        plt.close(fig)
        buf.seek(0)
        from PIL import Image
        w_px, h_px = Image.open(buf).size
        buf.seek(0)
        t = self.doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = (Inches(0.5), Inches(5.3), Inches(0.5))
        for c, w in zip(t.rows[0].cells, widths):
            c.width = w
        mid = t.rows[0].cells[1].paragraphs[0]
        mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width_in = min(w_px / 330.0, 5.2)
        mid.add_run().add_picture(buf, width=Inches(width_in))
        mid.paragraph_format.space_after = Pt(0)
        num = t.rows[0].cells[2].paragraphs[0]
        num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = num.add_run(f"({self.eq_n})"); r.font.name = FONT
        t.rows[0].cells[1].vertical_alignment = 1
        t.rows[0].cells[2].vertical_alignment = 1
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
        sp.paragraph_format.line_spacing = 0.5
        return self.eq_n

    # ---------------- figures ----------------
    def figure(self, path, caption, width=6.2):
        self.fig_n += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(3)
        para.add_run().add_picture(str(path), width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.paragraph_format.space_after = Pt(10)
        r = cap.add_run(f"Fig. {self.fig_n}. "); r.bold = True; r.font.size = Pt(9.5); r.font.name = FONT
        self._runs(cap, caption, size=9.5)
        return self.fig_n

    # ---------------- tables ----------------
    @staticmethod
    def _cell_border(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)
        for edge, sz in kw.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), "000000")
            borders.append(el)

    def table(self, header, rows, caption, widths=None, font=9, notes=None, align=None, bold_rows=(), shade_rows=()):
        self.tab_n += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.space_before = Pt(8)
        r = cap.add_run(f"Table {self.tab_n}. "); r.bold = True; r.font.size = Pt(9.5); r.font.name = FONT
        self._runs(cap, caption, size=9.5)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        align = align or ["l"] + ["c"] * (len(header) - 1)
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            para = c.paragraphs[0]
            para.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align[i]]
            para.paragraph_format.space_after = Pt(1)
            self._runs(para, f"**{htxt}**", size=font)
            self._cell_border(c, top=10, bottom=6)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                c = t.rows[ri].cells[ci]
                c.text = ""
                para = c.paragraphs[0]
                para.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align[ci]]
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.05
                txt = str(val)
                if (ri - 1) in bold_rows:
                    txt = f"**{txt}**" if not txt.startswith("**") else txt
                self._runs(para, txt, size=font)
                if ri == len(rows):
                    self._cell_border(c, bottom=10)
                if (ri - 1) in shade_rows:
                    tcPr = c._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EEF1F6")
                    tcPr.append(shd)
        if widths:
            for row in t.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Inches(w)
        if notes:
            n = self.doc.add_paragraph()
            n.paragraph_format.space_after = Pt(10)
            n.paragraph_format.space_before = Pt(2)
            self._runs(n, notes, size=8.5)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self.tab_n

    # ---------------- references ----------------
    def reference_list(self):
        for i, k in enumerate(self.cite_order, start=1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.35)
            para.paragraph_format.space_after = Pt(3)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._runs(para, f"[{i}] {self.refs[k]}", size=9.5)

    def save(self, path):
        self.doc.save(str(path))
