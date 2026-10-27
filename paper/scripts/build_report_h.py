"""Build the BCSE497J Project-I report in the layout of the reference PDF, with the wording taken from the humanized text.

Stage A  (log):    build the report once while logging every text segment (in document order).
Stage B  (align):  align the logged segments with the reference PDF text, and the reference PDF text with the humanized text,
                   so that every humanized word is assigned to exactly one segment (see align_text.py).
Stage C  (apply):  build again, this time writing the humanized wording into the same layout.
"""
import copy
import json
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import pandas as pd
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_ch1_4, report_ch5_9  # noqa: E402
from report_lib import Report, clear_paragraph, delete_paragraph, set_run_font  # noqa: E402
from report_refs import REFS  # noqa: E402
from report_tx import TX, tx_text  # noqa: E402

TEMPLATE = Path("/tmp/apo_report/template.docx")
WORK = Path("/tmp/apo_report")
FIG, RES = ROOT / "paper" / "figures", ROOT / "paper" / "results"

TITLE = "ADAPTIVE PORTFOLIO OPTIMIZATION IN DYNAMIC FINANCIAL MARKETS"
PROGRAMME = "Computer Science and Engineering (Artificial Intelligence and Machine Learning)"
STUDENTS = [("23BAI0002", "Yogdeep Benchimath"), ("23BAI0014", "Anushka Mishra"), ("23BAI0044", "Mehul Arora")]
GUIDE = "Dr. Arti Tiwari"
DESIGNATION = "Assistant Professor Sr. Grade 1"
HOD = "Dr. GAYATHRI P"

BT = json.load(open(RES / "backtest.json"))
M = json.load(open(RES / "matcher.json"))
ABL = json.load(open(RES / "ablation.json"))
TESTS = json.load(open(RES / "tests.json"))
abl = pd.DataFrame([dict(case=r["case"], cond=r["cond"], seed=r["seed"],
                         **{f"{s}_{m}": r["test"][s][k] for s in ("equal_weight", "mean_variance", "adaptive")
                            for m, k in (("sharpe", "sharpe_ratio"), ("vol", "annualized_volatility"), ("mdd", "max_drawdown"),
                                         ("ret", "total_return"))}) for r in ABL])


def small_img(path, max_w=1900):
    im = Image.open(path).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    out = path.with_suffix(".jpg")
    im.save(out, quality=88)
    return out


# ============================================================================ low-level template helpers
def first_rpr(p):
    for r in p.runs:
        if r._r.rPr is not None:
            return copy.deepcopy(r._r.rPr)
    return None


def _clean_rpr(rpr):
    for tag in ("w:highlight", "w:spacing", "w:w", "w:kern", "w:position", "w:fitText"):
        for h in rpr.findall(qn(tag)):
            rpr.remove(h)
    return rpr


def _style_run(r, rpr, size, bold, italic, highlight, underline):
    if rpr is not None:
        old = r._r.find(qn("w:rPr"))
        if old is not None:
            r._r.remove(old)
        r._r.insert(0, copy.deepcopy(rpr))
    set_run_font(r, size=size, bold=bold, italic=italic)
    if underline is not None:
        r.font.underline = underline
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_par(p, text, size=None, bold=None, italic=None, highlight=False, underline=None, route=True):
    """Replace the text of a paragraph (keeping paragraph properties and the first run's character formatting).
    Tabs are kept as tabs; every text piece goes through the text router."""
    rpr = first_rpr(p)
    if rpr is not None:
        rpr = _clean_rpr(rpr)
    clear_paragraph(p)
    last = None
    for part in re.split(r"(\t+)", text):
        if not part:
            continue
        r = p.add_run()
        if part.startswith("\t"):
            for _ in part:
                r.add_tab()
        else:
            r.add_text(tx_text(part) if route else part)
        _style_run(r, rpr, size, bold, italic, highlight, underline)
        last = r
    return last


def add_run(p, text, size=None, bold=None, italic=None, highlight=False, route=True):
    r = p.add_run(tx_text(text) if route else text)
    set_run_font(r, size=size, bold=bold, italic=italic)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def add_tab(p, size=12):
    r = p.add_run()
    r.add_tab()
    set_run_font(r, size=size)
    return r


W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


def strip_ids(el):
    for e in el.iter():
        for a in (W14 + "paraId", W14 + "textId"):
            if a in e.attrib:
                del e.attrib[a]


def clone_after(p, text=None, **kw):
    """Insert a copy of paragraph p right after it and optionally set its text."""
    new = copy.deepcopy(p._p)
    strip_ids(new)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    if text is not None:
        set_par(np_, text, **kw)
    return np_


def set_cell(cell, text, bold=None, size=None, highlight=False, route=True):
    p = cell.paragraphs[0]
    set_par(p, text, bold=bold, size=size, highlight=highlight, route=route)
    for extra in cell.paragraphs[1:]:
        delete_paragraph(extra)


def route_par(p, **kw):
    """Send the text of a template paragraph through the router, keeping its formatting."""
    if p.text.strip():
        set_par(p, p.text, **kw)


def add_pageref(p, bookmark, cached, size=11.5, bold=None):
    """Page number as a PAGEREF field (Word refreshes it with F9); the cached value is shown until then."""
    run = p.add_run()
    set_run_font(run, size=size, bold=bold)
    for kind, val in (("begin", None), ("instr", f" PAGEREF {bookmark} \\h "), ("separate", None), ("text", cached), ("end", None)):
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = val
        else:
            el = OxmlElement("w:t"); el.text = val
        run._r.append(el)
    return run


def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s; n -= v
    return out


# ============================================================================ front matter
def fill_front(doc, toc_rows, lof_rows, lot_rows):
    P = list(doc.paragraphs)
    T = doc.tables

    # ---------------- cover
    route_par(P[0])
    set_par(P[2], TITLE, size=18, bold=True)
    delete_paragraph(P[3])
    route_par(P[6])
    route_par(P[8])
    route_par(P[10])
    set_par(P[12], PROGRAMME, size=16, bold=True)
    route_par(P[14])
    for row, (reg, name) in zip(T[0].rows, STUDENTS):
        set_cell(row.cells[0], reg, bold=True)
        set_cell(row.cells[1], name.upper(), bold=True)
    delete_paragraph(P[16])
    route_par(P[18])
    P[18].paragraph_format.space_before = Pt(14)
    for i in (17, 20, 22):
        delete_paragraph(P[i])
    widen = (Inches(1.75), Inches(3.9))
    grid = T[0]._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widen):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    for row in T[0].rows:
        for c, w in zip(row.cells, widen):
            c.width = w
    T[0]._tbl.find(qn("w:tblPr")).find(qn("w:tblW")).set(qn("w:w"), str(int(sum(w.inches for w in widen) * 1440)))
    set_cell(T[1].rows[0].cells[0], GUIDE, bold=True)
    set_cell(T[1].rows[1].cells[0], DESIGNATION, highlight=True)
    set_cell(T[1].rows[2].cells[0], T[1].rows[2].cells[0].text)
    route_par(P[23])

    # ---------------- declaration
    route_par(P[25])
    p = P[27]
    clear_paragraph(p)
    for txt, it, bd in (("We hereby declare that the project entitled ", False, False), (TITLE, False, True),
                        (" submitted by us, for the award of the degree of ", False, False),
                        ("Bachelor of Technology in Computer Science and Engineering", True, False),
                        (f" to VIT is a record of bonafide work carried out by us under the supervision of {GUIDE}.", False, False)):
        add_run(p, txt, size=12, bold=bd or None, italic=it or None)
    set_par(P[28], "We further declare that the work reported in this project has not been submitted and will not be "
                   "submitted, either in part or in full, for the award of any other degree or diploma in this institute or "
                   "any other institute or university.", size=12)
    for i in (29, 30, 31, 32, 33):
        delete_paragraph(P[i])
    set_par(P[36], "Place\t: Vellore", size=12)
    P[36].paragraph_format.right_indent = Inches(0)
    clone_after(P[36], "Date\t:", size=12)
    set_par(P[37], "Signature of the Candidates", size=12, bold=True)
    P[37].paragraph_format.left_indent = Inches(3.1)
    last = P[37]
    for reg, name in STUDENTS:
        last = clone_after(last, f"{name} ({reg})", size=12, bold=False)
        last.paragraph_format.left_indent = Inches(3.1)
    for i in (38, 39, 40, 41, 42, 43):
        delete_paragraph(P[i])

    # ---------------- certificate
    route_par(P[45])
    p = P[47]
    clear_paragraph(p)
    names = ", ".join(f"{n} ({r})" for r, n in STUDENTS[:-1]) + f" and {STUDENTS[-1][1]} ({STUDENTS[-1][0]})"
    for txt, bd, it in (("This is to certify that the project entitled ", False, False), (TITLE, True, False),
                        (" submitted by ", False, False), (names, True, False),
                        (", School of Computer Science and Engineering, VIT, for the award of the degree of ", False, False),
                        ("Bachelor of Technology in Computer Science and Engineering", False, True),
                        (", is a record of bonafide work carried out by them under my supervision during Fall Semester "
                         "2026-2027, as per the VIT code of academic and research ethics.", False, False)):
        add_run(p, txt, size=12, bold=bd or None, italic=it or None)
    route_par(P[49])
    for i in (50, 51, 52, 53, 59, 60, 61, 62, 63, 64, 69, 70, 71):
        delete_paragraph(P[i])
    route_par(P[55])
    route_par(P[56])
    set_par(P[58], P[58].text, size=12, bold=True)
    clone_after(P[58], f"\t\t\t{GUIDE}", size=12, bold=False)
    set_par(P[66], P[66].text, size=12, bold=True)
    set_par(P[73], HOD, size=12, bold=True, highlight=True)
    clone_after(P[73], "HOD", size=12, bold=True, highlight=True)
    set_par(P[74], "B.Tech. " + PROGRAMME, size=12, bold=True)

    # ---------------- acknowledgements
    for i in (76, 77):
        delete_paragraph(P[i])
    set_par(P[75], P[75].text, size=14, bold=True, underline=WD_UNDERLINE.THICK)
    P[75].paragraph_format.page_break_before = True
    set_par(P[79], "We thank the management of Vellore Institute of Technology (VIT) for the opportunity to take up this "
                   "project and for the academic environment in which we could carry it out.", size=12)
    set_par(P[81], "We are grateful to Dr. Jaisankar N, Dean of the School of Computer Science and Engineering (SCOPE), for "
                   "the support and encouragement the school has given us during our studies.", size=12)
    p = P[83]
    clear_paragraph(p)
    add_run(p, f"We thank {HOD}, Head of the ", size=12)
    add_run(p, "B.Tech. " + PROGRAMME, size=12, bold=True)
    add_run(p, ", for the guidance and support given to us while we worked on this project.", size=12)
    p = P[85]
    clear_paragraph(p)
    add_run(p, f"We are especially thankful to our project guide, {GUIDE}, for the guidance, time and feedback that shaped "
               "this project.", size=12)
    set_par(P[87], "Thank you all for your contributions and support.", size=12)
    for i in (88, 89):
        delete_paragraph(P[i])
    set_par(P[91], STUDENTS[0][1], size=12, bold=True)
    P[91].paragraph_format.left_indent = Inches(3.1)
    last = P[91]
    for reg, name in STUDENTS[1:]:
        last = clone_after(last, name, size=12, bold=True)

    # ---------------- table of contents
    route_par(P[93])
    for i in (94, 95, 96, 97):
        delete_paragraph(P[i])
    toc = T[2]
    for c in toc.rows[0].cells:
        set_cell(c, c.text, bold=True)
    proto = {k: copy.deepcopy(toc.rows[i]._tr) for k, i in (("abs", 1), ("ch", 2), ("sec", 3), ("sub", 15))}
    for v in proto.values():
        strip_ids(v)
    for r in list(toc.rows)[1:]:
        toc._tbl.remove(r._tr)
    from docx.table import _Row
    for kind, a, b, bm, cached in toc_rows:
        tr = copy.deepcopy(proto[kind])
        toc._tbl.append(tr)
        row = _Row(tr, toc)
        for ci, (cell, txt) in enumerate(zip(row.cells, (a, b, None))):
            p = cell.paragraphs[0]
            sz = 12 if kind in ("abs", "ch") else 11.5
            bd = True if kind in ("abs", "ch") else False
            if ci == 2:
                clear_paragraph(p)
                add_pageref(p, bm, cached, size=sz, bold=bd)
            else:
                set_par(p, txt, size=sz, bold=bd, route=(ci == 1))
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- list of figures and list of tables
    for i in (100, 101, 102, 103, 104, 105):
        delete_paragraph(P[i])
    lof = T[3]
    lof_head = [c.text for c in lof.rows[0].cells]
    lot = copy.deepcopy(lof._tbl)
    strip_ids(lot)
    for i in (108, 109, 110):
        delete_paragraph(P[i])
    P[107]._p.addnext(lot)
    from docx.table import Table
    lot_t = Table(lot, P[107]._parent)
    set_par(P[99], "List of Figures", size=16, bold=True)
    P[99].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c, h in zip(lof.rows[0].cells, lof_head):
        set_cell(c, h, bold=True)
    fill_list_table(lof, lof_rows)
    route_par(P[107])
    P[107].paragraph_format.keep_with_next = True
    set_cell(lot_t.rows[0].cells[0], "Table No.", bold=True)
    set_cell(lot_t.rows[0].cells[1], lof_head[1], bold=True)
    set_cell(lot_t.rows[0].cells[2], lof_head[2], bold=True)
    fill_list_table(lot_t, lot_rows)
    prune_empty_rows(lof)
    prune_empty_rows(lot_t)
    prune_empty_rows(toc, col=1)

    # ---------------- abbreviations
    abbr = OrderedDict(sorted({
        "AD": "Adaptive (strategy)", "API": "Application Programming Interface", "CI": "Confidence Interval",
        "CNN": "Convolutional Neural Network", "CPU": "Central Processing Unit", "CSV": "Comma-Separated Values",
        "DDPG": "Deep Deterministic Policy Gradient", "DFD": "Data Flow Diagram", "DRL": "Deep Reinforcement Learning",
        "EM": "Expectation-Maximization", "ES": "Evolution Strategies", "ETF": "Exchange-Traded Fund",
        "EW": "Equal-Weight (strategy)", "FFN": "Feed-Forward Network", "FR": "Functional Requirement",
        "GAT": "Graph Attention Network", "GELU": "Gaussian Error Linear Unit", "GFC": "Global Financial Crisis",
        "GLD": "SPDR Gold Shares (ETF)", "GPU": "Graphics Processing Unit",
        "HDBSCAN": "Hierarchical Density-Based Spatial Clustering of Applications with Noise",
        "HMM": "Hidden Markov Model", "HTTP": "Hypertext Transfer Protocol", "IEEE": "Institute of Electrical and Electronics Engineers",
        "JSON": "JavaScript Object Notation", "k-NN": "k-Nearest Neighbors", "LOEO": "Leave-One-Event-Out",
        "LSTM": "Long Short-Term Memory", "MACD": "Moving Average Convergence Divergence", "MDD": "Maximum Drawdown",
        "MLP": "Multi-Layer Perceptron", "MV": "Mean-Variance (strategy)", "NFR": "Non-Functional Requirement",
        "PCA": "Principal Component Analysis", "PPO": "Proximal Policy Optimization", "QQQ": "Invesco QQQ Trust (Nasdaq-100 ETF)",
        "REST": "Representational State Transfer", "RL": "Reinforcement Learning", "RSI": "Relative Strength Index",
        "SDG": "Sustainable Development Goal", "SLSQP": "Sequential Least Squares Programming", "SMA": "Simple Moving Average",
        "SPY": "SPDR S&P 500 ETF Trust", "SR": "Sharpe Ratio", "t-SNE": "t-distributed Stochastic Neighbor Embedding",
        "TLT": "iShares 20+ Year Treasury Bond ETF", "TO": "Turnover", "TRL": "Technology Readiness Level",
        "UI": "User Interface", "VIX": "CBOE Volatility Index", "VUSTX": "Vanguard Long-Term Treasury Fund",
    }.items(), key=lambda kv: kv[0].lower()))
    route_par(P[112])
    proto_p = P[114]
    for i in range(115, 129):
        delete_paragraph(P[i])
    last = proto_p
    first = True
    for k, v in abbr.items():
        tgt = proto_p if first else clone_after(last)
        first = False
        clear_paragraph(tgt)
        add_run(tgt, k, size=12, route=False)
        add_tab(tgt)
        add_run(tgt, v, size=12)
        tgt.paragraph_format.space_after = Pt(2)
        tgt.paragraph_format.space_before = Pt(0)
        tgt.paragraph_format.left_indent = Inches(1.5)
        tgt.paragraph_format.first_line_indent = Inches(-1.5)
        for ts in list(tgt._p.pPr.findall(qn("w:tabs"))):
            tgt._p.pPr.remove(ts)
        tgt.paragraph_format.tab_stops.add_tab_stop(Inches(1.5))
        last = tgt

    # ---------------- symbols
    syms = [
        ("r_t", "Daily log return of an asset on day t"), ("P_t", "Adjusted closing price on day t"),
        ("σ_t", "20-day annualized rolling volatility"), ("d_t", "Drawdown from the 60-day rolling peak"),
        ("z_t", "Hidden regime state on day t (bull, bear or crisis)"), ("A", "State transition matrix of the HMM"),
        ("μ_j, σ_j", "Mean and standard deviation of the features in HMM state j"),
        ("X̃", "Length-invariant 40 by 3 shape of a price window"),
        ("c_τ, v_τ, g_τ", "Cumulative return, rolling volatility and running drawdown at point τ of a window"),
        ("d_model", "Hidden size of the Transformer (32)"), ("Q, K, V", "Query, key and value matrices of self-attention"),
        ("d_k", "Dimension of the attention keys"), ("f_θ, g_φ", "Encoder and decoder of the autoencoder"),
        ("L", "Reconstruction loss"), ("z", "16-dimensional embedding of a window"),
        ("ℓ", "Label of a corpus window (dot-com, GFC, COVID-19 or normal)"),
        ("k", "Number of nearest neighbors (8)"), ("s_ℓ", "In-class spacing of label ℓ"),
        ("m_ℓ", "Mean distance to the k nearest windows of label ℓ"), ("π_ℓ", "Gaussian-kernel similarity to label ℓ"),
        ("p_ℓ", "Local purity of label ℓ among the 15 nearest windows"), ("C_ℓ", "Match confidence for label ℓ"),
        ("w_k", "Portfolio weight vector at rebalance k"), ("s_k", "State vector given to the adaptive policy"),
        ("W₁, W₂, b₁, b₂", "Weights and biases of the policy network"), ("θ", "Flattened parameters of the policy"),
        ("ε_i", "Gaussian perturbation of the parameters in evolution strategies"),
        ("σ, α", "Perturbation scale (0.1) and step size (0.05) of evolution strategies"),
        ("F̂_i", "Rank-normalized reward of perturbation i"), ("R(θ)", "Training reward of the policy"),
        ("SR", "Sharpe ratio"), ("MDD", "Maximum drawdown"), ("TO_k", "Turnover at rebalance k"),
        ("V", "Portfolio value"), ("c", "Proportional transaction cost (5 basis points)"),
        ("r_f", "Annual risk-free rate (2 percent)"), ("Δ", "Difference in Sharpe ratio between two strategies"),
    ]
    for i in range(133, 184):
        delete_paragraph(P[i])
    route_par(P[130])
    p0 = P[132]
    clear_paragraph(p0)
    last = None
    for k, (s, mtxt) in enumerate(syms):
        tgt = p0 if k == 0 else clone_after(last)
        clear_paragraph(tgt)
        for tok in re.split(r"(_[^\s,]+)", s):
            if tok.startswith("_") and len(tok) > 1:
                add_run(tgt, tok[1:], size=12, route=False).font.subscript = True
            elif tok:
                add_run(tgt, tok, size=12, route=False)
        add_tab(tgt)
        add_run(tgt, mtxt, size=12)
        tgt.paragraph_format.left_indent = Inches(1.6)
        tgt.paragraph_format.first_line_indent = Inches(-1.6)
        tgt.paragraph_format.space_after = Pt(2)
        tgt.paragraph_format.space_before = Pt(0)
        tgt.paragraph_format.line_spacing = 1.0
        for ts in list(tgt._p.pPr.findall(qn("w:tabs"))):
            tgt._p.pPr.remove(ts)
        tgt.paragraph_format.tab_stops.add_tab_stop(Inches(1.6))
        last = tgt

    # ---------------- abstract
    for i in [185, 186, 188] + list(range(189, 225)):
        delete_paragraph(P[i])
    abstract = [
        "Portfolio managers allocate money with models that assume the market will behave tomorrow as it did in the sample "
        "used to fit them. Markets do not behave that way. The dot-com bust, the 2008 financial crisis and the COVID-19 crash "
        "each changed volatility, correlation and drawdown patterns within weeks, and static strategies such as mean-variance "
        "optimization react slowly. This project asks two practical questions. Does a chosen period of market data resemble "
        "one of those three crises, and does an allocation policy that knows the current market regime do better than simple "
        "benchmarks?",
        "We built a working web application, with a FastAPI back end and a React dashboard, that answers both questions on "
        "real daily prices from Yahoo Finance. A three-state Gaussian hidden Markov model labels each trading day as bull, "
        "bear or crisis without being given crash dates. For the comparison, any price window is turned into a 40-point, "
        "three-channel shape (cumulative return, short-term volatility and running drawdown), encoded into a 16-dimensional "
        "vector by a small Transformer encoder, and scored against a reference corpus of 965 labelled windows from the three "
        "crises and from calm periods. HDBSCAN is used to examine how the corpus clusters. The report shows the closest "
        "historical event, a confidence value, an overlay with the nearest real window, and a normal-regime baseline when no "
        "crisis matches. For allocation, we trained a small neural policy with evolution strategies and compared it with "
        "equal-weight and rolling mean-variance portfolios in walk-forward backtests that include a 5 basis point cost.",
        "The system passed 21 functional and unit tests, and the models were evaluated with leakage-controlled validation. The "
        "matcher separates crisis-like from normal periods, finding 48.8 percent of crisis windows with a 1.0 percent false "
        "alarm rate, and it flags crises it never saw in training in 53 to 93 percent of their windows. It cannot reliably say "
        "which crisis a period resembles (balanced accuracy 0.39 against a chance level of 0.25), and the trained Transformer "
        "did no better than an untrained one or seven summary statistics. The adaptive allocator did not beat mean-variance "
        "in three of four settings. Its one apparent win came from an invalid train and test split that produced a constant "
        "gold position, and every Sharpe ratio difference had a bootstrap interval that included zero. We report these results "
        "as they are, together with the diagnostics that exposed them. The project is at technology readiness level 3 and "
        "relates to Sustainable Development Goals 8 and 9.",
    ]
    set_par(P[184], "ABSTRACT", size=16, bold=True)
    P[184].paragraph_format.page_break_before = True
    P[184].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = P[187]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_par(p, abstract[0], size=12)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    last = p
    for t in abstract[1:]:
        last = clone_after(last, t, size=12)
    return P


def prune_empty_rows(tbl, col=1):
    """Drop contents rows whose text the humanized wording no longer contains."""
    for r in list(tbl.rows)[1:]:
        if not r.cells[col].text.strip():
            tbl._tbl.remove(r._tr)


def fill_list_table(tbl, rows):
    from docx.table import _Row
    proto = copy.deepcopy(tbl.rows[1]._tr)
    strip_ids(proto)
    for r in list(tbl.rows)[1:]:
        tbl._tbl.remove(r._tr)
    for no, title, bm, cached in rows:
        tr = copy.deepcopy(proto)
        tbl._tbl.append(tr)
        row = _Row(tr, tbl)
        for ci, cell in enumerate(row.cells):
            p = cell.paragraphs[0]
            if ci == 2:
                clear_paragraph(p)
                add_pageref(p, bm, cached, size=11)
            else:
                set_par(p, no if ci == 0 else title, size=11, route=(ci == 1))
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.keep_together = False
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cell in tbl.rows[0].cells:
        for r_ in cell.paragraphs[0].runs:
            r_.font.size = Pt(12)
    for tr_ in tbl._tbl.findall(qn("w:tr")):
        trPr = tr_.find(qn("w:trPr"))
        if trPr is not None:
            for h in trPr.findall(qn("w:trHeight")):
                trPr.remove(h)
    widths = (1250, 5700, 1250)
    for gc, w in zip(tbl._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)


PAGE_BORDER = (
    '<w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:offsetFrom="page">'
    '<w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="24" w:color="auto"/>'
    '<w:bottom w:val="single" w:sz="4" w:space="24" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="24" w:color="auto"/>'
    '</w:pgBorders>')


def add_page_borders(doc):
    from lxml import etree
    for sp in doc.element.body.iter(qn("w:sectPr")):
        if sp.find(qn("w:pgBorders")) is not None:
            continue
        mar = sp.find(qn("w:pgMar"))
        mar.addnext(etree.fromstring(PAGE_BORDER))


# ============================================================================ build once
def build(out_docx: Path, fig_labels=None, tab_labels=None, cached=None):
    cached = cached or {}
    doc = Document(str(TEMPLATE))
    P0 = list(doc.paragraphs)
    for p in P0[226:]:
        delete_paragraph(p)
    fig_labels = defaultdict(lambda: "?.?", fig_labels or {})
    tab_labels = defaultdict(lambda: "?.?", tab_labels or {})
    R = Report(doc, REFS, fig_labels, tab_labels)
    D = dict(FIG=FIG, BT=BT, M=M, ABL=ABL, TESTS=TESTS, ROOT=ROOT, abl=abl, small_img=small_img)
    report_ch1_4.write(R, D)
    report_ch5_9.write(R, D)

    sects = doc.sections
    final_sect = sects[-1]._sectPr
    front = sects[1]._sectPr
    for tag in ("w:pgSz", "w:pgMar", "w:pgNumType"):
        for el in final_sect.findall(qn(tag)):
            final_sect.remove(el)
    for tag in ("w:pgSz", "w:pgMar"):
        final_sect.append(copy.deepcopy(front.find(qn(tag))))
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal"); pg.set(qn("w:start"), "1")
    final_sect.append(pg)
    for el in final_sect.findall(qn("w:headerReference")) + final_sect.findall(qn("w:footerReference")):
        final_sect.remove(el)
    cols = final_sect.find(qn("w:cols"))
    if cols is not None:
        final_sect.remove(cols); final_sect.append(cols)

    sec = doc.sections[-1]
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    set_run_font(run, size=11)
    for tag, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        elif text.strip() == "PAGE":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        else:
            el = OxmlElement("w:t"); el.text = text
        run._r.append(el)

    # ---- contents tables (page numbers are PAGEREF fields with cached values)
    toc_rows, lof_rows, lot_rows = [], [], []
    first_bm = R.headings[0][2]
    toc_rows.append(("abs", "", "Abstract", cached.get("__abstract_bm__", first_bm), cached.get("ABSTRACT", "")))
    for lvl, text, bm in R.headings:
        pgn = cached.get(text, "")
        if lvl == 0:
            num, title = text.split("|", 1)
            if num == "A":
                toc_rows.append(("ch", "", "APPENDIX A: SAMPLE CODE", bm, pgn))
            else:
                toc_rows.append(("ch", f"{num}.", title, bm, pgn))
        elif lvl == 1:
            toc_rows.append(("sec", "", text, bm, pgn))
        elif lvl == 2:
            if text.startswith("A."):
                continue
            toc_rows.append(("sub", "", "     " + text, bm, pgn))
    clean = lambda s: re.sub(r"\{(?:fig|tab):[a-z_0-9]+\}", "", re.sub(r"\{cite:[^}]+\}", "", s)).replace("**", "").replace("*", "")
    for label, cap, bm in R.figures:
        lof_rows.append((label, clean(cap).rstrip("."), bm, cached.get(f"Figure {label}", "")))
    for label, cap, bm in R.tables:
        lot_rows.append((label, clean(cap).rstrip("."), bm, cached.get(f"Table {label}", "")))
    toc_rows[0] = ("abs", "", "Abstract", "_RefApoAbstract", cached.get("ABSTRACT", ""))
    TX.front_start = TX.counter
    fill_front(doc, toc_rows, lof_rows, lot_rows)
    ensure_abstract_bookmark(doc)
    add_page_borders(doc)
    doc.save(str(out_docx))
    return R, D


def ensure_abstract_bookmark(doc):
    """Bookmark on the ABSTRACT heading paragraph of the template (paragraph 184 of the original)."""
    for p in doc.paragraphs:
        if p.text.strip() == "ABSTRACT":
            start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), "999"); start.set(qn("w:name"), "_RefApoAbstract")
            end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), "999")
            if p._p.pPr is not None:
                p._p.pPr.addnext(start)
            else:
                p._p.insert(0, start)
            p._p.append(end)
            return "_RefApoAbstract"
    raise RuntimeError("abstract heading not found")
