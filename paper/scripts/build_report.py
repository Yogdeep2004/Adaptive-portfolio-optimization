"""Build the BCSE497J Project-I report on the VIT template.
Pass 1 builds the document and renders it to PDF through Pages, the page numbers of every heading, figure and table are
read from that PDF, and pass 2 rebuilds the document with those numbers in the contents tables.

Run:  /opt/anaconda3/bin/python paper/scripts/build_report.py
"""
import copy
import json
import pickle
import re
import subprocess
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import pandas as pd
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_ch1_4, report_ch5_9  # noqa: E402
from report_lib import Report, clear_paragraph, delete_paragraph, set_run_font  # noqa: E402
from report_refs import REFS  # noqa: E402

TEMPLATE = Path("/tmp/apo_report/template.docx")
WORK = Path("/tmp/apo_report")
OUT_DOCX = Path("/Users/yogdeepbenchimath/Downloads/BCSE497J_Project_I_Report_Adaptive_Portfolio_Optimizer.docx")
OUT_PDF = OUT_DOCX.with_suffix(".pdf")
FIG, RES = ROOT / "paper" / "figures", ROOT / "paper" / "results"

TITLE = "ADAPTIVE PORTFOLIO OPTIMIZATION IN DYNAMIC FINANCIAL MARKETS"
PROGRAMME = "Computer Science and Engineering (Artificial Intelligence and Machine Learning)"
STUDENTS = [("23BAI0002", "Yogdeep Benchimath"), ("23BAI0014", "Anushka Mishra"), ("23BAI0044", "Mehul Arora")]
GUIDE = "Dr. Arti Tiwari"

BT = json.load(open(RES / "backtest.json"))
M = json.load(open(RES / "matcher.json"))
ABL = json.load(open(RES / "ablation.json"))
TESTS = json.load(open(RES / "tests.json"))
abl = pd.DataFrame([dict(case=r["case"], cond=r["cond"], seed=r["seed"],
                         **{f"{s}_{m}": r["test"][s][k] for s in ("equal_weight", "mean_variance", "adaptive")
                            for m, k in (("sharpe", "sharpe_ratio"), ("vol", "annualized_volatility"), ("mdd", "max_drawdown"),
                                         ("ret", "total_return"))}) for r in ABL])


def small_img(path, max_w=1900):
    im = Image.open(path).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    out = path.with_suffix(".jpg")
    im.save(out, quality=88)
    return out


# ============================================================================ low-level template helpers
def first_rpr(p):
    for r in p.runs:
        if r._r.rPr is not None:
            return copy.deepcopy(r._r.rPr)
    return None


def set_par(p, text, size=None, bold=None, italic=None, highlight=False, underline=None):
    """Replace the text of a paragraph, keeping the paragraph properties and the first run's character formatting."""
    rpr = first_rpr(p)
    clear_paragraph(p)
    r = p.add_run(text)
    if rpr is not None:
        old = r._r.find(qn("w:rPr"))
        if old is not None:
            r._r.remove(old)
        for tag in ("w:highlight", "w:spacing", "w:w", "w:kern", "w:position", "w:fitText"):
            for h in rpr.findall(qn(tag)):
                rpr.remove(h)
        r._r.insert(0, rpr)
    set_run_font(r, size=size, bold=bold, italic=italic)
    if underline is not None:
        r.font.underline = underline
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def add_run(p, text, size=None, bold=None, italic=None, highlight=False):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


def strip_ids(el):
    for e in el.iter():
        for a in (W14 + "paraId", W14 + "textId"):
            if a in e.attrib:
                del e.attrib[a]


def clone_after(p, text=None, **kw):
    """Insert a copy of paragraph p right after it and optionally set its text."""
    new = copy.deepcopy(p._p)
    strip_ids(new)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    if text is not None:
        set_par(np_, text, **kw)
    return np_


def strip_highlight(p):
    for r in p.runs:
        r.font.highlight_color = None


def set_cell(cell, text, bold=None, size=None, highlight=False):
    p = cell.paragraphs[0]
    set_par(p, text, bold=bold, size=size, highlight=highlight)
    for extra in cell.paragraphs[1:]:
        delete_paragraph(extra)


def to_pdf(docx, pdf):
    subprocess.run([str(WORK / "topdf.sh"), str(docx), str(pdf)], check=True, timeout=300)


def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s; n -= v
    return out


# ============================================================================ front matter
def fill_front(doc, toc_rows, lof_rows, lot_rows):
    P = list(doc.paragraphs)
    T = doc.tables

    # ---------------- cover
    set_par(P[2], TITLE, size=18, bold=True)
    delete_paragraph(P[3])
    set_par(P[12], PROGRAMME, size=16, bold=True)
    for row, (reg, name) in zip(T[0].rows, STUDENTS):
        set_cell(row.cells[0], reg, bold=True)
        set_cell(row.cells[1], name.upper(), bold=True)
    delete_paragraph(P[16])
    P[18].paragraph_format.space_before = Pt(14)
    for i in (17, 20, 22):
        delete_paragraph(P[i])
    widen = (Inches(1.75), Inches(3.9))
    grid = T[0]._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widen):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    for row in T[0].rows:
        for c, w in zip(row.cells, widen):
            c.width = w
    T[0]._tbl.find(qn("w:tblPr")).find(qn("w:tblW")).set(qn("w:w"), str(int(sum(w.inches for w in widen) * 1440)))
    set_cell(T[1].rows[0].cells[0], GUIDE, bold=True)
    set_cell(T[1].rows[1].cells[0], "[Designation of the guide]", highlight=True)

    # ---------------- declaration
    p = P[27]
    clear_paragraph(p)
    for txt, it, bd in (("We hereby declare that the project entitled ", False, False), (TITLE, False, True),
                        (" submitted by us, for the award of the degree of ", False, False),
                        ("Bachelor of Technology in Computer Science and Engineering", True, False),
                        (f" to VIT is a record of bonafide work carried out by us under the supervision of {GUIDE}.", False, False)):
        add_run(p, txt, size=12, bold=bd or None, italic=it or None)
    set_par(P[28], "We further declare that the work reported in this project has not been submitted and will not be "
                   "submitted, either in part or in full, for the award of any other degree or diploma in this institute or "
                   "any other institute or university.", size=12)
    for i in (29, 30, 31, 32, 33):
        delete_paragraph(P[i])
    set_par(P[36], "Place\t: Vellore", size=12)
    P[36].paragraph_format.right_indent = Inches(0)
    clone_after(P[36], "Date\t:", size=12)
    set_par(P[37], "Signature of the Candidates", size=12, bold=True)
    P[37].paragraph_format.left_indent = Inches(3.1)
    last = P[37]
    for reg, name in STUDENTS:
        last = clone_after(last, f"{name} ({reg})", size=12, bold=False)
        last.paragraph_format.left_indent = Inches(3.1)
    for i in (38, 39, 40, 41, 42, 43):
        delete_paragraph(P[i])

    # ---------------- certificate
    p = P[47]
    clear_paragraph(p)
    names = ", ".join(f"{n} ({r})" for r, n in STUDENTS[:-1]) + f" and {STUDENTS[-1][1]} ({STUDENTS[-1][0]})"
    for txt, bd, it in (("This is to certify that the project entitled ", False, False), (TITLE, True, False),
                        (" submitted by ", False, False), (names, True, False),
                        (", School of Computer Science and Engineering, VIT, for the award of the degree of ", False, False),
                        ("Bachelor of Technology in Computer Science and Engineering", False, True),
                        (", is a record of bonafide work carried out by them under my supervision during Fall Semester "
                         "2026-2027, as per the VIT code of academic and research ethics.", False, False)):
        add_run(p, txt, size=12, bold=bd or None, italic=it or None)
    for i in (50, 51, 52, 53, 59, 60, 61, 62, 63, 64, 69, 70, 71):
        delete_paragraph(P[i])
    clone_after(P[58], f"\t\t\t{GUIDE}", size=12, bold=False)
    set_par(P[73], "[Name of the Head of the Department]", size=12, bold=True, highlight=True)
    set_par(P[74], "B.Tech. " + PROGRAMME, size=12, bold=True)

    # ---------------- acknowledgements
    for i in (76, 77):
        delete_paragraph(P[i])
    set_par(P[79], "We thank the management of Vellore Institute of Technology (VIT) for the opportunity to take up this "
                   "project and for the academic environment in which we could carry it out.", size=12)
    set_par(P[81], "We are grateful to Dr. Jaisankar N, Dean of the School of Computer Science and Engineering (SCOPE), for "
                   "the support and encouragement the school has given us during our studies.", size=12)
    p = P[83]
    clear_paragraph(p)
    add_run(p, "We thank ", size=12)
    add_run(p, "[Name of the Head of the Department]", size=12, highlight=True)
    add_run(p, ", Head of the ", size=12)
    add_run(p, "[Department Name]", size=12, highlight=True)
    add_run(p, ", for the guidance and support given to us while we worked on this project.", size=12)
    p = P[85]
    clear_paragraph(p)
    add_run(p, f"We are especially thankful to our project guide, {GUIDE}, for the guidance, time and feedback that shaped "
               "this project.", size=12)
    set_par(P[87], "Thank you all for your contributions and support.", size=12)
    for i in (88, 89):
        delete_paragraph(P[i])
    set_par(P[91], STUDENTS[0][1], size=12, bold=True)
    P[91].paragraph_format.left_indent = Inches(3.1)
    last = P[91]
    for reg, name in STUDENTS[1:]:
        last = clone_after(last, name, size=12, bold=True)

    # ---------------- table of contents
    for i in (94, 95, 96, 97):
        delete_paragraph(P[i])
    toc = T[2]
    proto = {k: copy.deepcopy(toc.rows[i]._tr) for k, i in (("abs", 1), ("ch", 2), ("sec", 3), ("sub", 15))}
    for v in proto.values():
        strip_ids(v)
    for r in list(toc.rows)[1:]:
        toc._tbl.remove(r._tr)
    from docx.table import _Row
    for kind, a, b, c in toc_rows:
        tr = copy.deepcopy(proto[kind])
        toc._tbl.append(tr)
        row = _Row(tr, toc)
        for cell, txt in zip(row.cells, (a, b, c)):
            p = cell.paragraphs[0]
            set_par(p, txt)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            for r_ in p.runs:
                r_.font.size = Pt(12 if kind in ("abs", "ch") else 11.5)
                r_.bold = True if kind in ("abs", "ch") else False
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- list of figures and list of tables
    for i in (100, 101, 102, 103, 104, 105):
        delete_paragraph(P[i])
    lof = T[3]
    fill_list_table(lof, lof_rows)
    # list of tables: clone the figure list table and move it under the heading
    lot = copy.deepcopy(lof._tbl)
    strip_ids(lot)
    for i in (108, 109, 110):
        delete_paragraph(P[i])
    P[107].paragraph_format.keep_with_next = True
    P[107]._p.addnext(lot)
    from docx.table import Table
    lot_t = Table(lot, P[107]._parent)
    set_cell(lot_t.rows[0].cells[0], "Table No.", bold=True)
    fill_list_table(lot_t, lot_rows)
    set_par(P[99], "List of Figures", size=16, bold=True)
    P[99].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- abbreviations
    abbr = OrderedDict(sorted({
        "AD": "Adaptive (strategy)", "API": "Application Programming Interface", "CI": "Confidence Interval",
        "CNN": "Convolutional Neural Network", "CPU": "Central Processing Unit", "CSV": "Comma-Separated Values",
        "DDPG": "Deep Deterministic Policy Gradient", "DFD": "Data Flow Diagram", "DRL": "Deep Reinforcement Learning",
        "EM": "Expectation-Maximization", "ES": "Evolution Strategies", "ETF": "Exchange-Traded Fund",
        "EW": "Equal-Weight (strategy)", "FFN": "Feed-Forward Network", "FR": "Functional Requirement",
        "GAT": "Graph Attention Network", "GELU": "Gaussian Error Linear Unit", "GFC": "Global Financial Crisis",
        "GLD": "SPDR Gold Shares (ETF)", "GPU": "Graphics Processing Unit",
        "HDBSCAN": "Hierarchical Density-Based Spatial Clustering of Applications with Noise",
        "HMM": "Hidden Markov Model", "HTTP": "Hypertext Transfer Protocol", "IEEE": "Institute of Electrical and Electronics Engineers",
        "JSON": "JavaScript Object Notation", "k-NN": "k-Nearest Neighbors", "LOEO": "Leave-One-Event-Out",
        "LSTM": "Long Short-Term Memory", "MACD": "Moving Average Convergence Divergence", "MDD": "Maximum Drawdown",
        "MLP": "Multi-Layer Perceptron", "MV": "Mean-Variance (strategy)", "NFR": "Non-Functional Requirement",
        "PCA": "Principal Component Analysis", "PPO": "Proximal Policy Optimization", "QQQ": "Invesco QQQ Trust (Nasdaq-100 ETF)",
        "REST": "Representational State Transfer", "RL": "Reinforcement Learning", "RSI": "Relative Strength Index",
        "SDG": "Sustainable Development Goal", "SLSQP": "Sequential Least Squares Programming", "SMA": "Simple Moving Average",
        "SPY": "SPDR S&P 500 ETF Trust", "SR": "Sharpe Ratio", "t-SNE": "t-distributed Stochastic Neighbor Embedding",
        "TLT": "iShares 20+ Year Treasury Bond ETF", "TO": "Turnover", "TRL": "Technology Readiness Level",
        "UI": "User Interface", "VIX": "CBOE Volatility Index", "VUSTX": "Vanguard Long-Term Treasury Fund",
    }.items(), key=lambda kv: kv[0].lower()))
    proto_p = P[114]
    for i in range(115, 129):
        delete_paragraph(P[i])
    last = proto_p
    first = True
    for k, v in abbr.items():
        if first:
            set_par(proto_p, f"{k}\t{v}", size=12); first = False
        else:
            last = clone_after(last, f"{k}\t{v}", size=12)
    # the template's abbreviation tab stop is narrow, so widen it for long entries
    for p in doc.paragraphs:
        if "\t" in p.text and p.text.split("\t")[0] in abbr and p.text.split("\t")[1] == abbr[p.text.split("\t")[0]]:
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.left_indent = Inches(1.5)
            p.paragraph_format.first_line_indent = Inches(-1.5)
            p.paragraph_format.tab_stops.clear_all() if hasattr(p.paragraph_format.tab_stops, "clear_all") else None
            p.paragraph_format.tab_stops.add_tab_stop(Inches(1.5))

    # ---------------- symbols
    syms = [
        ("r_t", "Daily log return of an asset on day t"), ("P_t", "Adjusted closing price on day t"),
        ("σ_t", "20-day annualized rolling volatility"), ("d_t", "Drawdown from the 60-day rolling peak"),
        ("z_t", "Hidden regime state on day t (bull, bear or crisis)"), ("A", "State transition matrix of the HMM"),
        ("μ_j, σ_j", "Mean and standard deviation of the features in HMM state j"),
        ("X̃", "Length-invariant 40 by 3 shape of a price window"),
        ("c_τ, v_τ, g_τ", "Cumulative return, rolling volatility and running drawdown at point τ of a window"),
        ("d_model", "Hidden size of the Transformer (32)"), ("Q, K, V", "Query, key and value matrices of self-attention"),
        ("d_k", "Dimension of the attention keys"), ("f_θ, g_φ", "Encoder and decoder of the autoencoder"),
        ("L", "Reconstruction loss"), ("z", "16-dimensional embedding of a window"),
        ("ℓ", "Label of a corpus window (dot-com, GFC, COVID-19 or normal)"),
        ("k", "Number of nearest neighbors (8)"), ("s_ℓ", "In-class spacing of label ℓ"),
        ("m_ℓ", "Mean distance to the k nearest windows of label ℓ"), ("π_ℓ", "Gaussian-kernel similarity to label ℓ"),
        ("p_ℓ", "Local purity of label ℓ among the 15 nearest windows"), ("C_ℓ", "Match confidence for label ℓ"),
        ("w_k", "Portfolio weight vector at rebalance k"), ("s_k", "State vector given to the adaptive policy"),
        ("W₁, W₂, b₁, b₂", "Weights and biases of the policy network"), ("θ", "Flattened parameters of the policy"),
        ("ε_i", "Gaussian perturbation of the parameters in evolution strategies"),
        ("σ, α", "Perturbation scale (0.1) and step size (0.05) of evolution strategies"),
        ("F̂_i", "Rank-normalized reward of perturbation i"), ("R(θ)", "Training reward of the policy"),
        ("SR", "Sharpe ratio"), ("MDD", "Maximum drawdown"), ("TO_k", "Turnover at rebalance k"),
        ("V", "Portfolio value"), ("c", "Proportional transaction cost (5 basis points)"),
        ("r_f", "Annual risk-free rate (2 percent)"), ("Δ", "Difference in Sharpe ratio between two strategies"),
    ]
    for i in range(133, 184):
        delete_paragraph(P[i])
    p0 = P[132]
    clear_paragraph(p0)
    pf = p0.paragraph_format
    last = None
    for k, (s, mtxt) in enumerate(syms):
        tgt = p0 if k == 0 else clone_after(last)
        clear_paragraph(tgt)
        for tok in re.split(r"(_[^\s,]+)", s):
            if tok.startswith("_") and len(tok) > 1:
                add_run(tgt, tok[1:], size=12).font.subscript = True
            elif tok:
                add_run(tgt, tok, size=12)
        add_run(tgt, f"\t{mtxt}", size=12)
        tgt.paragraph_format.left_indent = Inches(1.6)
        tgt.paragraph_format.first_line_indent = Inches(-1.6)
        tgt.paragraph_format.space_after = Pt(2)
        tgt.paragraph_format.space_before = Pt(0)
        tgt.paragraph_format.line_spacing = 1.0
        for ts in list(tgt._p.pPr.findall(qn("w:tabs"))):
            tgt._p.pPr.remove(ts)
        tgt.paragraph_format.tab_stops.add_tab_stop(Inches(1.6))
        last = tgt

    # ---------------- abstract
    for i in [185, 186, 188] + list(range(189, 225)):
        delete_paragraph(P[i])
    abstract = [
        "Portfolio managers allocate money with models that assume the market will behave tomorrow as it did in the sample "
        "used to fit them. Markets do not behave that way. The dot-com bust, the 2008 financial crisis and the COVID-19 crash "
        "each changed volatility, correlation and drawdown patterns within weeks, and static strategies such as mean-variance "
        "optimization react slowly. This project asks two practical questions. Does a chosen period of market data resemble "
        "one of those three crises, and does an allocation policy that knows the current market regime do better than simple "
        "benchmarks?",
        "We built a working web application, with a FastAPI back end and a React dashboard, that answers both questions on "
        "real daily prices from Yahoo Finance. A three-state Gaussian hidden Markov model labels each trading day as bull, "
        "bear or crisis without being given crash dates. For the comparison, any price window is turned into a 40-point, "
        "three-channel shape (cumulative return, short-term volatility and running drawdown), encoded into a 16-dimensional "
        "vector by a small Transformer encoder, and scored against a reference corpus of 965 labelled windows from the three "
        "crises and from calm periods. HDBSCAN is used to examine how the corpus clusters. The report shows the closest "
        "historical event, a confidence value, an overlay with the nearest real window, and a normal-regime baseline when no "
        "crisis matches. For allocation, we trained a small neural policy with evolution strategies and compared it with "
        "equal-weight and rolling mean-variance portfolios in walk-forward backtests that include a 5 basis point cost.",
        "The system passed 21 functional and unit tests, and the models were evaluated with leakage-controlled validation. The "
        "matcher separates crisis-like from normal periods, finding 48.8 percent of crisis windows with a 1.0 percent false "
        "alarm rate, and it flags crises it never saw in training in 53 to 93 percent of their windows. It cannot reliably say "
        "which crisis a period resembles (balanced accuracy 0.39 against a chance level of 0.25), and the trained Transformer "
        "did no better than an untrained one or seven summary statistics. The adaptive allocator did not beat mean-variance "
        "in three of four settings. Its one apparent win came from an invalid train and test split that produced a constant "
        "gold position, and every Sharpe ratio difference had a bootstrap interval that included zero. We report these results "
        "as they are, together with the diagnostics that exposed them. The project is at technology readiness level 3 and "
        "relates to Sustainable Development Goals 8 and 9.",
    ]
    p = P[187]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_par(p, abstract[0], size=12)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    last = p
    for t in abstract[1:]:
        last = clone_after(last, t, size=12)
    set_par(P[184], "ABSTRACT", size=16, bold=True)
    P[184].paragraph_format.page_break_before = True
    P[184].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return P


def fill_list_table(tbl, rows):
    from docx.table import _Row
    proto = copy.deepcopy(tbl.rows[1]._tr)
    strip_ids(proto)
    for r in list(tbl.rows)[1:]:
        tbl._tbl.remove(r._tr)
    for no, title, page in rows:
        tr = copy.deepcopy(proto)
        tbl._tbl.append(tr)
        row = _Row(tr, tbl)
        for cell, txt in zip(row.cells, (no, title, page)):
            p = cell.paragraphs[0]
            set_par(p, txt, size=11)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.keep_together = False
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cell in tbl.rows[0].cells:
        for r_ in cell.paragraphs[0].runs:
            r_.font.size = Pt(12)
    for tr_ in tbl._tbl.findall(qn("w:tr")):
        trPr = tr_.find(qn("w:trPr"))
        if trPr is not None:
            for h in trPr.findall(qn("w:trHeight")):
                trPr.remove(h)
    widths = (1250, 5700, 1250)
    for gc, w in zip(tbl._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)


# ============================================================================ build once
def build(pages: dict | None, out_docx: Path, fig_labels=None, tab_labels=None):
    doc = Document(str(TEMPLATE))
    P0 = list(doc.paragraphs)
    # remove the template's body (chapter 1 placeholder onward) but keep the section break paragraph 225
    body_paras = P0[226:]
    for p in body_paras:
        delete_paragraph(p)
    # the body-level sectPr now closes the main section: give it the main section's margins and numbering
    body = doc.element.body
    final = body.find(qn("w:sectPr"))

    fig_labels = defaultdict(lambda: "?.?", fig_labels or {})
    tab_labels = defaultdict(lambda: "?.?", tab_labels or {})
    R = Report(doc, REFS, fig_labels, tab_labels)
    D = dict(FIG=FIG, BT=BT, M=M, ABL=ABL, TESTS=TESTS, ROOT=ROOT, abl=abl, small_img=small_img)
    report_ch1_4.write(R, D)
    report_ch5_9.write(R, D)

    # ---- section properties of the body (arabic numbers restart at 1; margins as in the template's main section)
    sects = doc.sections
    final_sect = sects[-1]._sectPr
    # copy pgSz/pgMar from the front-matter section, then set arabic numbering
    front = sects[1]._sectPr
    for tag in ("w:pgSz", "w:pgMar", "w:pgNumType"):
        for el in final_sect.findall(qn(tag)):
            final_sect.remove(el)
    for tag in ("w:pgSz", "w:pgMar"):
        final_sect.append(copy.deepcopy(front.find(qn(tag))))
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal"); pg.set(qn("w:start"), "1")
    final_sect.append(pg)
    for el in final_sect.findall(qn("w:headerReference")) + final_sect.findall(qn("w:footerReference")):
        final_sect.remove(el)
    # schema order: references, then pgSz, pgMar, pgNumType, cols
    cols = final_sect.find(qn("w:cols"))
    if cols is not None:
        final_sect.remove(cols); final_sect.append(cols)

    # ---- own footer for the body with a plain decimal page number
    sec = doc.sections[-1]
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    set_run_font(run, size=11)
    for tag, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        elif text.strip() == "PAGE":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        else:
            el = OxmlElement("w:t"); el.text = text
        run._r.append(el)

    # ---- contents tables
    toc_rows, lof_rows, lot_rows = [], [], []
    pg_of = (pages or {})
    toc_rows.append(("abs", "", "Abstract", pg_of.get("ABSTRACT", "")))
    for lvl, text in R.headings:
        key = text
        if lvl == 0:
            num, title = text.split("|", 1)
            if num == "A":
                toc_rows.append(("ch", "", "APPENDIX A: SAMPLE CODE", pg_of.get(key, "")))
            else:
                toc_rows.append(("ch", f"{num}.", title, pg_of.get(key, "")))
        elif lvl == 1:
            toc_rows.append(("sec", "", text, pg_of.get(key, "")))
        elif lvl == 2:
            if text.startswith("A."):
                continue
            toc_rows.append(("sub", "", "     " + text, pg_of.get(key, "")))
    for label, cap in R.figures:
        lof_rows.append((label, cap.rstrip("."), pg_of.get(f"Figure {label}", "")))
    for label, cap in R.tables:
        lot_rows.append((label, cap.rstrip("."), pg_of.get(f"Table {label}", "")))
    # strip inline markup left in captions
    clean = lambda s: re.sub(r"\{(?:fig|tab):[a-z_0-9]+\}", "", re.sub(r"\{cite:[^}]+\}", "", s)).replace("**", "").replace("*", "")
    lof_rows = [(a, clean(b), c) for a, b, c in lof_rows]
    lot_rows = [(a, clean(b), c) for a, b, c in lot_rows]
    fill_front(doc, toc_rows, lof_rows, lot_rows)
    doc.save(str(out_docx))
    return R, D


def detect_pages(pdf_path, R):
    """Return {key: displayed page number string} from the rendered PDF."""
    d = pymupdf.open(str(pdf_path))
    texts = [re.sub(r"[ \t]+", " ", pg.get_text()) for pg in d]
    # first body page: contains the line 'CHAPTER 1' followed by INTRODUCTION
    first_body = next(i for i, t in enumerate(texts) if re.search(r"^CHAPTER 1\s*\nINTRODUCTION", t, re.M))
    abs_page = next(i for i, t in enumerate(texts[:first_body]) if re.search(r"^ABSTRACT\s*$", t, re.M))
    pages = {"ABSTRACT": roman(abs_page)}   # cover is physical page 0 and unnumbered, so physical index i shows roman(i)

    def disp(i):
        return str(i - first_body + 1)

    for lvl, text in R.headings:
        if lvl == 0:
            num, title = text.split("|", 1)
            pat = r"^CHAPTER %s\s*\n" % num if num != "A" else r"^APPENDIX A: SAMPLE CODE"
        else:
            pat = r"^" + re.escape(text.strip()) + r"\s*$"
        for i in range(first_body, len(texts)):
            if re.search(pat, texts[i], re.M):
                pages[text] = disp(i)
                break
    for label, cap in R.figures:
        start = re.escape(f"Figure {label} ")
        for i in range(first_body, len(texts)):
            if re.search(r"^" + start, texts[i], re.M):
                pages[f"Figure {label}"] = disp(i); break
    for label, cap in R.tables:
        start = re.escape(f"Table {label} ")
        for i in range(first_body, len(texts)):
            if re.search(r"^" + start, texts[i], re.M):
                pages[f"Table {label}"] = disp(i); break
    return pages, len(d), first_body


if __name__ == "__main__":
    p1 = WORK / "pass1.docx"
    R, D = build(None, p1)
    labels = (dict(R.fig_labels), dict(R.tab_labels))
    print("pass 1 built; figures", len(R.figures), "tables", len(R.tables), "refs", len(R.cite_order))
    # second dry build so that cross-references resolve to the final labels
    R, D = build(None, p1, *labels)
    to_pdf(p1, WORK / "pass1.pdf")
    pages, n, fb = detect_pages(WORK / "pass1.pdf", R)
    print("pass 1 pages:", n, "first body page index", fb, "headings located", len(pages))
    for attempt in range(4):
        R, D = build(pages, OUT_DOCX, *labels)
        to_pdf(OUT_DOCX, OUT_PDF)
        pages2, n2, fb2 = detect_pages(OUT_PDF, R)
        diff = {k: (pages.get(k), v) for k, v in pages2.items() if pages.get(k) != v}
        print("pass", attempt + 2, "| pages:", n2, "| page-number changes:", diff)
        if not diff:
            break
        pages = pages2
    json.dump(pages2, open(WORK / "pages_final.json", "w"), indent=1)
