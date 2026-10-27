"""Build the report in the layout of the reference PDF with the humanized wording.

  /opt/anaconda3/bin/python paper/scripts/make_h_report.py

Inputs : the reference PDF (layout oracle) and the humanized DOCX (wording), both from ~/Downloads.
Output : ~/Downloads/BCSE497J_Project_I_Report_Humanized_Formatted.docx (plus a clean, un-humanized twin for comparison).
"""
import difflib
import json
import re
import sys
from collections import defaultdict, deque
from pathlib import Path

import pymupdf
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_report_h import build, roman, WORK  # noqa: E402
from report_tx import TX  # noqa: E402

DL = Path("/Users/yogdeepbenchimath/Downloads")
REF_PDF = DL / "Project Report Adaptive Portfolio Optimizer.pdf"
H_DOCX = DL / "h_Project Report Adaptive Portfolio Optimizer.docx"
OUT = DL / "BCSE497J_Project_I_Report_Humanized_Formatted.docx"
OUT_CLEAN = DL / "BCSE497J_Project_I_Report_Clean_Formatted.docx"


# ---------------------------------------------------------------------------------------------- token utilities
def key(t):
    k = re.sub(r"[^a-z0-9]", "", t.lower())
    return k or t


def merge_tokens(tokens, vocab):
    """Repair artefacts of PDF line breaks and of the PDF-to-DOCX conversion (split hyphens, split words)."""
    out, i, n = [], 0, len(tokens)
    while i < n:
        t = tokens[i]
        if t.endswith("-") and len(t) > 1 and t[-2].isalnum() and i + 1 < n and tokens[i + 1][:1].isalnum():
            out.append(t + tokens[i + 1]); i += 2; continue
        if t == "-" and out and out[-1][-1:].isalpha() and i + 1 < n and tokens[i + 1][:1].islower():
            out[-1] = out[-1] + "-" + tokens[i + 1]; i += 2; continue
        if out and t.startswith("-") and len(t) > 1:
            prev = out[-1]
            if ((t[1].isalpha() and prev[-1:].isalpha())
                    or (re.fullmatch(r"-\d\d?[,.;:]?", t) and prev.isalpha() and prev.isupper())
                    or (re.fullmatch(r"-\d{4}\S*", t) and re.fullmatch(r"\d{4}", prev))):
                out[-1] = prev + t; i += 1; continue
        out.append(t); i += 1
    res, i = [], 0
    while i < len(out):
        a = out[i]
        if i + 1 < len(out) and a.isalpha():
            m = re.fullmatch(r"([A-Za-z]+)([^\w]*)", out[i + 1])
            if m:
                core = m.group(1)
                if (a + core).lower() in vocab and (a.lower() not in vocab or core.lower() not in vocab):
                    res.append(a + out[i + 1]); i += 2; continue
        res.append(a); i += 1
    return res


def load_pages(vocab):
    ref = pymupdf.open(str(REF_PDF))
    R = [merge_tokens(pg.get_text().split(), vocab) for pg in ref]
    d = Document(str(H_DOCX))
    H = [merge_tokens(p.text.split(), vocab) for p in d.paragraphs]
    assert len(R) == len(H), (len(R), len(H))
    return R, H


# ---------------------------------------------------------------------------------------------- alignment
def align(blocks):
    """blocks: logged segments in document order. Returns {bid: new text}."""
    vocab = set()
    for b in blocks:
        for w in re.findall(r"[A-Za-z]+", b["text"]):
            vocab.add(w.lower())
    R_pages, H_pages = load_pages(vocab)
    R = [t for pg in R_pages for t in pg]
    page_of = [p for p, pg in enumerate(R_pages) for _ in pg]
    page_start = {}
    for i, p in enumerate(page_of):
        page_start.setdefault(p, i)
    rkeys = [key(t) for t in R]

    # --- my non-table blocks vs the reference text
    flat, owners = [], []
    for b in blocks:
        if b["ctx"] is None:
            for ti, t in enumerate(b["text"].split()):
                flat.append(key(t)); owners.append((b["bid"], ti))
    sm = difflib.SequenceMatcher(None, flat, rkeys, autojunk=False)
    owner = {}
    for a, r, size in sm.get_matching_blocks():
        for k in range(size):
            owner[r + k] = owners[a + k]
    matched_blocks = sum(1 for _ in set(o[0] for o in owner.values()))
    print(f"  my blocks matched in reference text: {matched_blocks} of {sum(1 for b in blocks if b['ctx'] is None)}"
          f" non-table blocks; {len(owner)} of {len(flat)} tokens")

    # --- table cells: region between neighbouring matched blocks, tokens assigned cell by cell
    seq = blocks
    first_match_after = [None] * (len(seq) + 1)
    nxt = None
    first_match_after[len(seq)] = None
    block_rmin = {}
    for r, (bid, ti) in owner.items():
        block_rmin[bid] = min(r, block_rmin.get(bid, 1 << 60))
    block_rmax = {}
    for r, (bid, ti) in owner.items():
        block_rmax[bid] = max(r, block_rmax.get(bid, -1))
    bid_pos = {b["bid"]: i for i, b in enumerate(seq)}
    for i in range(len(seq) - 1, -1, -1):
        b = seq[i]
        if b["ctx"] is None and b["bid"] in block_rmin:
            nxt = block_rmin[b["bid"]]
        first_match_after[i] = nxt
    last_r = -1
    tables = defaultdict(list)
    table_span = {}
    for i, b in enumerate(seq):
        if b["ctx"] is None:
            if b["bid"] in block_rmax:
                last_r = max(last_r, block_rmax[b["bid"]])
        else:
            tid = b["ctx"][0]
            if tid not in table_span:
                table_span[tid] = (last_r + 1, first_match_after[i])
            tables[tid].append(b)
    assigned = total = 0
    for tid, cells in tables.items():
        lo, hi = table_span[tid]
        hi = len(R) if hi is None else hi
        rows = defaultdict(lambda: defaultdict(deque))
        for b in cells:
            _, ri, ci = b["ctx"]
            for ti, t in enumerate(b["text"].split()):
                rows[ri][ci].append((b["bid"], ti, key(t)))
        order = sorted(rows)
        total += sum(len(q) for r_ in rows.values() for q in r_.values())
        rr, miss = 0, 0
        for r in range(lo, hi):
            k = rkeys[r]
            while rr < len(order) and all(len(q) == 0 for q in rows[order[rr]].values()):
                rr += 1
            if rr >= len(order):
                break
            hit = None
            for ci in sorted(rows[order[rr]]):
                q = rows[order[rr]][ci]
                if q and q[0][2] == k:
                    hit = q; break
            if hit is None and miss >= 4 and rr + 1 < len(order):
                for ci in sorted(rows[order[rr + 1]]):
                    q = rows[order[rr + 1]][ci]
                    if q and q[0][2] == k:
                        hit = q; rr += 1; break
            if hit is not None:
                bid, ti, _ = hit.popleft()
                owner[r] = (bid, ti); miss = 0; assigned += 1
            else:
                miss += 1
    print(f"  table cell tokens assigned: {assigned} of {total}")

    # --- reference text -> humanized text, page by page; every humanized token gets one reference owner
    att = defaultdict(list)       # reference token index -> [(humanized token, is_equal)]
    pre = defaultdict(list)
    for p in range(len(R_pages)):
        a0 = page_start.get(p, 0)
        rp, hp = R_pages[p], H_pages[p]
        sm = difflib.SequenceMatcher(None, [key(t) for t in rp], [key(t) for t in hp], autojunk=False)
        for tag, a, b, c, d in sm.get_opcodes():
            if tag == "equal":
                for k in range(b - a):
                    att[a0 + a + k].append((hp[c + k], True))
            elif tag == "replace":
                n, m = b - a, d - c
                if m >= n:
                    for k in range(n - 1):
                        att[a0 + a + k].append((hp[c + k], False))
                    att[a0 + b - 1].extend((hp[c + k], False) for k in range(n - 1, m))
                else:
                    for k in range(m):
                        att[a0 + a + k].append((hp[c + k], False))
            elif tag == "insert":
                tgt = a0 + a - 1 if a > 0 else a0
                (att if a > 0 else pre)[tgt].extend((hp[c + k], False) for k in range(d - c))

    new = defaultdict(list)
    last_owner, pending = None, []
    align.dropped = []
    for r in range(len(R)):
        own = owner.get(r)
        items = pre.get(r, []) + att.get(r, [])
        if own is not None:
            new[own[0]].extend(pending); pending = []
            new[own[0]].extend((t, own[1], eq) for t, eq in items)
            last_owner = own[0]
        else:
            align.dropped.extend((r, page_of[r], t) for t, eq in items if eq)
            edits = [t for t, eq in items if not eq]
            if edits:
                if r - page_start[page_of[r]] < 4 or last_owner is None:
                    pending.extend((t, None, False) for t in edits)
                else:
                    new[last_owner].extend((t, None, False) for t in edits)
    n_h = sum(len(p) for p in H_pages)
    n_used = sum(len(v) for v in new.values())
    new = {bid: list(v) for bid, v in new.items()}
    print(f"  humanized tokens: {n_h}; placed into text blocks: {n_used}")
    return new, R_pages, owner


# ---------------------------------------------------------------------------------------------- page numbers from the reference PDF
def ref_page_numbers(R_obj):
    d = pymupdf.open(str(REF_PDF))
    texts = [re.sub(r"[ \t]+", " ", pg.get_text()) for pg in d]
    first_body = next(i for i, t in enumerate(texts) if re.search(r"^CHAPTER 1\s*\n", t, re.M))
    abs_page = next(i for i, t in enumerate(texts[:first_body]) if re.search(r"^ABSTRACT\s*$", t, re.M))
    pages = {"ABSTRACT": roman(abs_page)}

    def disp(i):
        return str(i - first_body + 1)

    for lvl, text, bm in R_obj.headings:
        if lvl == 0:
            num, title = text.split("|", 1)
            pat = r"^CHAPTER %s\s*\n" % num if num != "A" else r"^APPENDIX A: SAMPLE CODE"
        else:
            pat = r"^" + re.escape(text.strip()) + r"\s*$"
        for i in range(first_body, len(texts)):
            if re.search(pat, texts[i], re.M):
                pages[text] = disp(i); break

    def cap_start(cap):
        c = re.sub(r"\{(?:fig|tab):[a-z_0-9]+\}", "", re.sub(r"\{cite:[^}]+\}", "", cap)).replace("**", "").replace("*", "")
        return re.escape(" ".join(c.split()[:3]))

    for kind, items in (("Figure", R_obj.figures), ("Table", R_obj.tables)):
        for label, cap, bm in items:
            pat = r"^" + re.escape(f"{kind} {label}") + r"\s+" + cap_start(cap)
            for i in range(first_body, len(texts)):
                if re.search(pat, texts[i], re.M):
                    pages[f"{kind} {label}"] = disp(i); break
    return pages


if __name__ == "__main__":
    # 0) dry build to learn the figure and table numbers
    TX.reset("off")
    R0, _ = build(WORK / "dry.docx")
    labels = (dict(R0.fig_labels), dict(R0.tab_labels))
    cached = ref_page_numbers(R0)
    print("cached page numbers found for", len(cached), "entries")

    # clean twin (no humanized wording), with the user's own edits
    TX.reset("off")
    build(OUT_CLEAN, *labels, cached=cached)

    # A) log every text segment
    TX.reset("log")
    build(WORK / "log.docx", *labels, cached=cached)
    blocks = list(TX.blocks)
    front = [b for b in blocks if b["bid"] >= TX.front_start]
    body = [b for b in blocks if b["bid"] < TX.front_start]
    doc_order = front + body
    print("logged segments:", len(blocks), "(front matter", len(front), ", body", len(body), ")")

    # B) align
    new, R_pages, owner = align(doc_order)
    json.dump({str(k): " ".join(w for w, _, _ in v) for k, v in new.items()}, open(WORK / "new_text.json", "w"), indent=0)

    # C) apply
    TX.reset("apply", new)
    build(OUT, *labels, cached=cached)
    print("saved", OUT)
